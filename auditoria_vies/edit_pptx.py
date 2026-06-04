# Corrige lacunas da apresentação e insere os slides da auditoria empírica
# (com as figuras). Gera backup apresentacao_backup_original.pptx.
from pptx import Presentation
from pptx.util import Pt, Inches, Emu
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN
import shutil, copy, json
from pathlib import Path

RES = Path("auditoria_vies/resultados")
shutil.copy("apresentacao.pptx", "apresentacao_backup_original.pptx")
prs = Presentation("apresentacao.pptx")
S = prs.slides

ACCENT = RGBColor(0x1F, 0x4E, 0x79)   # azul-escuro para destaques

def body_shape(slide):
    """Retorna o placeholder de corpo (idx 1) ou maior textbox."""
    for sh in slide.placeholders:
        if sh.placeholder_format.idx == 1:
            return sh
    # fallback: maior textbox
    cands=[sh for sh in slide.shapes if sh.has_text_frame]
    return max(cands, key=lambda s: s.height) if cands else None

def set_bullets(tf, items, size=20, clear=True):
    if clear:
        tf.clear()
    for i, it in enumerate(items):
        p = tf.paragraphs[0] if (i == 0 and clear) else tf.add_paragraph()
        if isinstance(it, tuple):
            text, lvl = it
        else:
            text, lvl = it, 0
        p.level = lvl
        run = p.add_run(); run.text = text
        run.font.size = Pt(size - 4*lvl)
        run.font.name = "Calibri"
    return tf

# FIX slide 2 — lista completa de autores (8)
autores = ["Artemio Costa Canossa","Bruna de Carvalho Regis","Fernando Bandeira Soares",
           "Kelly Cristina Maximo de Almeida","Maria Luiza Ribeiro de Araujo",
           "Mario Henrique Gimenes Santana","Rafael Oliveira da Cruz","Valdina Alves Cogue"]
for sh in S[1].placeholders:
    if sh.placeholder_format.idx == 14 and "Canossa" in sh.text_frame.text:
        tf = sh.text_frame; tf.clear()
        for i,a in enumerate(autores):
            p = tf.paragraphs[0] if i==0 else tf.add_paragraph()
            r=p.add_run(); r.text=a; r.font.size=Pt(18); r.font.name="Calibri"

# FIX slide 9 — Metodologia coerente com a auditoria empirica
b9 = None
for sh in S[8].shapes:
    if sh.has_text_frame and "qualitativa" in sh.text_frame.text.lower():
        b9 = sh; break
if b9:
    set_bullets(b9.text_frame, [
        "Abordagem qualitativa: análise bibliográfica e documental (eixo jurídico)",
        "Abordagem quantitativa: auditoria empírica de viés (eixo de Ciência de Dados)",
        ("Auditoria de modelo pré-treinado (ArcFace) sobre bases públicas — sem treino do zero", 1),
        ("Sem uso de imagens de pessoas reais desaparecidas (conformidade ética/LGPD)", 1),
    ], size=20)

# FIX títulos de seção desalinhados + slide 17 (trabalhos futuros reais)
def set_title(slide, text):
    for sh in slide.placeholders:
        if sh.placeholder_format.idx == 0:
            sh.text_frame.clear(); r=sh.text_frame.paragraphs[0].add_run()
            r.text=text; return

set_title(S[12], "Resultados — Marco Jurídico Brasileiro")   # era "Resultados"
set_title(S[14], "Experiências Internacionais")              # era "Conclusão"

# slide 17 -> trabalhos futuros reais (alinhados ao TCC 8.1)
for sh in S[16].shapes:
    if sh.has_text_frame and "Integração" in sh.text_frame.text:
        set_bullets(sh.text_frame, [
            "Avaliar base de faces representativa da população brasileira",
            "Estender a auditoria à identificação 1:N e a múltiplos modelos",
            "Investigar envelhecimento facial (casos de crianças)",
            "Propor protocolo de homologação de sistemas biométricos humanitários",
        ], size=20)

# FIX slide 19 — encerramento
for sh in S[18].placeholders:
    if sh.placeholder_format.idx == 14:
        tf=sh.text_frame; tf.clear()
        linhas=["Obrigado!","Análise de Viabilidade Técnica e Jurídica da IA",
                "na Identificação de Pessoas Desaparecidas",
                "UNIVESP — Bacharelado em Ciência de Dados — 2026"]
        for i,l in enumerate(linhas):
            p=tf.paragraphs[0] if i==0 else tf.add_paragraph()
            r=p.add_run(); r.text=l; r.font.size=Pt(28 if i==0 else 18)
            r.font.name="Calibri"; p.alignment=PP_ALIGN.CENTER
        break

# NOVOS SLIDES — Auditoria Empírica (com figuras reais)
lfw=json.load(open(RES/"lfw_metrics.json")); ff=json.load(open(RES/"fairface_summary.json"))
def br(x,n): return (f"%.{n}f"%x).replace(".",",")
auc=br(lfw["AUC"],4); eer=br(lfw["EER"]*100,2); acc=br(lfw["best_accuracy"]*100,2)
disp=br(ff["disparity_max_min"],2)

LAY_OBJ = prs.slide_layouts[2]   # OBJECT (titulo+corpo)
LAY_TITLE = prs.slide_layouts[6] # TITLE_ONLY

def new_obj_slide(title, bullets, size=20):
    sl = S.add_slide(LAY_OBJ)
    set_title(sl, title)
    bd = body_shape(sl)
    set_bullets(bd.text_frame, bullets, size=size)
    return sl

def add_textbox(sl, left, top, width, height, lines):
    tb = sl.shapes.add_textbox(Inches(left),Inches(top),Inches(width),Inches(height))
    tf = tb.text_frame; tf.word_wrap=True
    for i,(txt,sz,bold,color) in enumerate(lines):
        p=tf.paragraphs[0] if i==0 else tf.add_paragraph()
        r=p.add_run(); r.text=txt; r.font.size=Pt(sz); r.font.bold=bold
        r.font.name="Calibri"
        if color: r.font.color.rgb=color
    return tb

# Slide N1 — Método da auditoria
s_metodo = new_obj_slide("Auditoria Empírica de Viés — Contribuição de Ciência de Dados", [
    "Auditoria de modelo pré-treinado ArcFace (ResNet-50), sem treino do zero",
    "Bases públicas: LFW (6.000 pares) e FairFace (10.954 imagens rotuladas)",
    "Verificação facial 1:1 por similaridade de cosseno entre embeddings",
    "Limiar único + métricas: AUC, EER, FMR, FNMR e teste qui-quadrado",
    "Reprodutível em Python (ONNX Runtime, scikit-learn) — sem GPU",
], size=20)

# Slide N2 — Resultados desempenho (LFW) + Figura 1
s_lfw = S.add_slide(LAY_TITLE)
set_title(s_lfw, "Resultados — Desempenho de Verificação (LFW)")
add_textbox(s_lfw, 0.6, 1.4, 8.8, 1.0, [
    (f"AUC = {auc}    •    EER = {eer}%    •    Acurácia = {acc}%", 24, True, ACCENT),
    ("Alto desempenho em condições controladas — confirma a literatura", 16, False, None),
])
s_lfw.shapes.add_picture(str(RES/"fig1_lfw_roc.png"), Inches(0.7), Inches(2.7), width=Inches(8.6))

# Slide N3 — Resultados viés (FairFace) + Figura 2
s_ff = S.add_slide(LAY_TITLE)
set_title(s_ff, "Resultados — Viés Demográfico (FairFace)")
add_textbox(s_ff, 0.6, 1.35, 8.8, 1.5, [
    (f"Disparidade de FMR = {disp}× entre grupos  (χ² = 80,1; p < 0,001)", 22, True, ACCENT),
    ("Grupo negro ≈ 2× a taxa de falso positivo do grupo branco", 16, False, None),
    ("Mulheres não brancas as mais afetadas — replica o padrão Gender Shades", 16, False, None),
])
s_ff.shapes.add_picture(str(RES/"fig2_fairface_fmr.png"), Inches(1.3), Inches(2.9), width=Inches(7.4))

# REORDENAR — inserir os 3 novos slides após o slide 12 (Contexto brasileiro)
sldIdLst = prs.slides._sldIdLst
ids = list(sldIdLst)
new_ids = [ids[-3], ids[-2], ids[-1]]   # os 3 recém-criados (no fim)
for nid in new_ids: sldIdLst.remove(nid)
# posição de inserção: logo após o 12º slide "Contexto brasileiro" (encerra bloco técnico/viés)
insert_at = 12
for k,nid in enumerate(new_ids):
    sldIdLst.insert(insert_at+k, nid)

prs.save("apresentacao.pptx")
print("OK — apresentacao.pptx revisada. Slides agora:", len(prs.slides.__iter__.__self__._sldIdLst))
print(f"LFW AUC={auc} EER={eer}% acc={acc}% | FairFace disparidade={disp}x")
